import io
import fitz  # PyMuPDF
import docx
from docx.opc.exceptions import PackageNotFoundError
from fastapi import UploadFile, HTTPException


async def parse_file(file: UploadFile) -> str:
    filename = (file.filename or "").lower()
    content = await file.read()

    # ---------------------------------------------------------
    # PDF EXTRACTION
    # ---------------------------------------------------------
    if filename.endswith(".pdf"):
        try:
            pdf_doc = fitz.open(stream=content, filetype="pdf")

            text_parts = []

            for page_number in range(pdf_doc.page_count):
                page = pdf_doc.load_page(page_number)
                page_text = page.get_text("text")

                if isinstance(page_text, str) and page_text.strip():
                    text_parts.append(page_text)

            extracted_text = "\n".join(text_parts).strip()

            if extracted_text:
                return extracted_text

            raise HTTPException(
                status_code=400,
                detail="Resume is incompatible or scanned. Please upload a text-based PDF or DOCX file."
            )

        except HTTPException:
            raise

        except Exception:
            raise HTTPException(
                status_code=400,
                detail="Invalid or corrupted PDF file."
            )

    # ---------------------------------------------------------
    # DOCX EXTRACTION (paragraphs + tables)
    # ---------------------------------------------------------
    if filename.endswith(".docx"):
        try:
            word_doc = docx.Document(io.BytesIO(content))

            # Extract paragraphs
            paragraph_text = "\n".join(
                p.text for p in word_doc.paragraphs
            )

            # Extract tables
            table_text_parts = []

            for table in word_doc.tables:
                for row in table.rows:
                    row_cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    row_text = " | ".join(
                        c for c in row_cells if c
                    )

                    if row_text:
                        table_text_parts.append(row_text)

            table_text = "\n".join(table_text_parts)

            combined_text = (
                paragraph_text + "\n" + table_text
            ).strip()

            if combined_text:
                return combined_text

            raise HTTPException(
                status_code=400,
                detail="Resume content could not be extracted from DOCX."
            )

        except PackageNotFoundError:
            raise HTTPException(
                status_code=400,
                detail="Invalid or corrupted DOCX file."
            )

        except HTTPException:
            raise

        except Exception:
            raise HTTPException(
                status_code=400,
                detail="Unable to read DOCX file."
            )

    # ---------------------------------------------------------
    # UNSUPPORTED FILE TYPES
    # ---------------------------------------------------------
    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Please upload PDF or DOCX."
    )